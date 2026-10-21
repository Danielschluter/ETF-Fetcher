#!/usr/bin/env python3
"""Scrape the latest holdings of every US-listed ETF.

Holdings come from the fund issuers' own websites, which publish them daily,
plus a catch-all source for every ETF without a dedicated issuer scraper:

  ishares      iShares holdings CSV (BlackRock)
  ssga         SPDR daily holdings workbook (State Street)
  invesco      Invesco holdings API
  first_trust  First Trust holdings page
  global_x     Global X daily holdings CSV
  vanguard     Vanguard holdings API (month-end holdings)
  proshares    ProShares fund page
  innovator    Innovator daily holdings CSV, one file for all its funds
  pacer        Pacer daily holdings CSV
  schwab       Schwab all-holdings pages
  vaneck       VanEck daily holdings workbook
  ark          ARK daily holdings CSV
  direxion     Direxion daily holdings CSV
  dimensional  Dimensional daily holdings CSV
  avantis      Avantis total-holdings page (American Century)
  american_century  American Century total-holdings page
  fidelity     Fidelity research holdings page (month-end holdings, weights only)
  goldman_sachs  Goldman Sachs fund data service (GraphQL)
  zacks       Zacks ETF holdings page: used for all other ETFs, and as the
               fallback whenever an issuer source fails. Zacks turns away plain
               HTTP clients, so its pages are loaded in headless Chrome (via
               Playwright). Zacks does not say which date its holdings are as
               of, so as_of_date is left blank.

The ETF universe is every security flagged as an ETF in the Nasdaq Trader
symbol directory, which covers all US exchanges (NYSE, Arca, Nasdaq, Cboe...).

Outputs (in --output-dir)
  holdings.csv     one row per position, all ETFs combined
  etf_summary.csv  one row per ETF: source used, as-of date, number of holdings,
                   total weight, or why no holdings were found

Examples
  python etf_holdings_scraper.py
  python etf_holdings_scraper.py --tickers SPY IVV QQQ
  python etf_holdings_scraper.py --sources ishares ssga --no-fallback
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import gzip
import hashlib
import html
import io
import json
import logging
import os
import queue
import re
import statistics
import sys
import threading
import time
import warnings
from collections import Counter
from concurrent.futures import Future, ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable, Sequence, TypeVar
from urllib.parse import urlsplit

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

log = logging.getLogger("etf_holdings")
T = TypeVar("T")

# SSGA workbooks have no default style, which openpyxl warns about on every file.
warnings.filterwarnings("ignore", message="Workbook contains no default style", module="openpyxl")

NASDAQ_SYMBOL_DIRECTORY = "https://www.nasdaqtrader.com/dynamic/SymDir/nasdaqtraded.txt"

DEFAULT_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36"
)
HOLDINGS_TTL = 6 * 3600  # issuers publish once per business day
REFERENCE_TTL = 24 * 3600

# Websites whose robots.txt asks for more seconds between requests than --rate allows.
CRAWL_DELAYS = {"www.vaneck.com": 25}

EXCHANGES = {
    "A": "NYSE American",
    "N": "NYSE",
    "P": "NYSE Arca",
    "Q": "Nasdaq",
    "V": "IEX",
    "Z": "Cboe BZX",
}

HOLDING_FIELDS = [
    "holding_name",
    "ticker",
    "cusip",
    "isin",
    "sedol",
    "other_id",
    "asset_class",
    "sector",
    "sector_source",
    "country",
    "currency",
    "shares",
    "price",
    "market_value",
    "notional_value",
    "weight_pct",
    "value_weight_pct",
    "maturity_date",
    "coupon_rate",
]
HOLDING_COLUMNS = ["etf_ticker", "etf_name", "source", "as_of_date", *HOLDING_FIELDS]
NUMERIC_FIELDS = {"shares", "price", "market_value", "notional_value", "weight_pct",
                  "value_weight_pct", "coupon_rate"}

SUMMARY_COLUMNS = [
    "ticker",
    "name",
    "exchange",
    "status",
    "source",
    "as_of_date",
    "holdings_count",
    "weight_sum_pct",
    "source_url",
    "detail",
]


# --------------------------------------------------------------------------- HTTP


class NotAvailable(Exception):
    """A source has no holdings for this fund; the next source should be tried."""


class NotFound(NotAvailable):
    """HTTP 404/410."""


class Blocked(NotAvailable):
    """A website answered with a bot-check page instead of the content."""


# Pages that bot-protection services (Imperva, Cloudflare) serve instead of the content.
_BOT_CHECK = re.compile(
    r"Pardon Our Interruption|<title>\s*Just a moment\.\.\.\s*</title>", re.IGNORECASE)


def looks_like_bot_check(resp: requests.Response) -> bool:
    if "html" not in resp.headers.get("Content-Type", "") or len(resp.content) > 100_000:
        return False
    return bool(_BOT_CHECK.search(resp.content.decode("utf-8", errors="replace")))


class RateLimiter:
    """Spaces requests at least 1/rate seconds apart across all threads."""

    def __init__(self, rate: float):
        self.interval = 1.0 / rate
        self.lock = threading.Lock()
        self.next_slot = time.monotonic()

    def wait(self) -> None:
        with self.lock:
            now = time.monotonic()
            slot = max(now, self.next_slot)
            self.next_slot = slot + self.interval
        if slot > now:
            time.sleep(slot - now)

    def pause(self, seconds: float) -> bool:
        """Hold every request for a while. False if they were already held about as long."""
        with self.lock:
            until = time.monotonic() + seconds
            if until <= self.next_slot + 1:
                return False
            self.next_slot = until
            return True


class HttpClient:
    """GET with a per-host rate limit and retries; one requests.Session per thread.

    Bot protection (on Zacks, for one) lets most requests through but answers
    some with a check page. Then all requests to that site pause and the page
    is retried; a site that keeps blocking is skipped for the rest of the run.
    """

    RETRY_STATUSES = {429, 500, 502, 503, 504}
    BOT_CHECK_PAUSES = (30, 60, 120)  # seconds
    GIVE_UP_AFTER = 3  # pages in a row still blocked after every pause

    def __init__(self, user_agent: str, rate: float, timeout: float = 60, retries: int = 3):
        self.headers = {
            "User-Agent": user_agent,
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "en-US,en;q=0.9",
        }
        self.rate = rate
        self.timeout = timeout
        self.retries = retries
        self.local = threading.local()
        self.lock = threading.Lock()
        self.limiters: dict[str, RateLimiter] = {}
        self.blocked_streak: Counter = Counter()
        self.given_up: set[str] = set()

    def _session(self) -> requests.Session:
        session = getattr(self.local, "session", None)
        if session is None:
            session = requests.Session()
            session.headers.update(self.headers)
            self.local.session = session
        return session

    def _limiter(self, host: str) -> RateLimiter:
        with self.lock:
            if host not in self.limiters:
                rate = min(self.rate, 1 / CRAWL_DELAYS[host]) if host in CRAWL_DELAYS else self.rate
                self.limiters[host] = RateLimiter(rate)
            return self.limiters[host]

    def get(self, url: str) -> bytes:
        return self._fetch(url)

    def post_json(self, url: str, payload: dict, headers: dict | None = None) -> bytes:
        return self._fetch(url, payload, headers)

    def _fetch(self, url: str, payload: dict | None = None,
               headers: dict | None = None) -> bytes:
        host = urlsplit(url).hostname or ""
        limiter = self._limiter(host)
        for pause in (*self.BOT_CHECK_PAUSES, None):
            if host in self.given_up:
                raise Blocked(f"skipped: {host} kept answering with a bot check")
            resp = self._send_with_retries(url, limiter, payload, headers)
            if not looks_like_bot_check(resp):
                with self.lock:
                    self.blocked_streak[host] = 0
                return resp.content
            if pause is not None and limiter.pause(pause):
                log.warning("%s answered with a bot check; pausing requests to it for %ds",
                            host, pause)
        with self.lock:
            self.blocked_streak[host] += 1
            if self.blocked_streak[host] >= self.GIVE_UP_AFTER and host not in self.given_up:
                self.given_up.add(host)
                log.error("%s keeps answering with a bot check; skipping it for the rest of "
                          "this run. Run again later to fill in the gaps.", host)
        raise Blocked(f"{host} answered with a bot check")

    def _send_with_retries(self, url: str, limiter: RateLimiter, payload: dict | None,
                           headers: dict | None) -> requests.Response:
        for attempt in range(self.retries + 1):
            limiter.wait()
            delay = 2.0 ** (attempt + 1)
            try:
                if payload is None:
                    resp = self._session().get(url, timeout=self.timeout)
                else:
                    resp = self._session().post(url, json=payload, headers=headers,
                                                timeout=self.timeout)
            except requests.RequestException as exc:
                error: Exception = exc
            else:
                if resp.status_code in (404, 410):
                    raise NotFound(f"HTTP {resp.status_code}")
                if resp.status_code not in self.RETRY_STATUSES:
                    resp.raise_for_status()
                    return resp
                error = requests.HTTPError(f"HTTP {resp.status_code}", response=resp)
                retry_after = resp.headers.get("Retry-After", "")
                if retry_after.isdigit():
                    delay = max(delay, float(retry_after))
            if attempt == self.retries:
                raise error
            log.debug("Retrying %s in %.0fs after: %s", url, delay, error)
            time.sleep(delay)
        raise AssertionError("unreachable")


class DiskCache:
    """Gzipped file cache. Writes are atomic so an interrupted run never leaves a bad entry."""

    def __init__(self, root: Path | None):
        self.root = root

    def _path(self, url: str) -> Path:
        host = urlsplit(url).hostname or "other"
        return self.root / host / f"{hashlib.sha1(url.encode()).hexdigest()}.gz"

    def read(self, url: str, max_age: float) -> bytes | None:
        if self.root is None:
            return None
        path = self._path(url)
        try:
            if time.time() - path.stat().st_mtime > max_age:
                return None
            return gzip.decompress(path.read_bytes())
        except (OSError, EOFError):
            return None

    def write(self, url: str, data: bytes) -> None:
        if self.root is None:
            return
        path = self._path(url)
        path.parent.mkdir(parents=True, exist_ok=True)
        tmp = path.with_name(f"{path.name}.{os.getpid()}.{threading.get_ident()}.tmp")
        tmp.write_bytes(gzip.compress(data, compresslevel=6))
        os.replace(tmp, path)

    def delete(self, url: str) -> None:
        if self.root is not None:
            self._path(url).unlink(missing_ok=True)


class SetupError(RuntimeError):
    """Something the run needs could not be started."""


class BrowserPool:
    """Loads pages in headless Chrome, for websites that turn scripts away but serve browsers.

    Playwright objects can only be used by the thread that created them, so each
    worker thread runs its own browser and takes URLs from a shared queue.
    """

    BOT_CHECK_WAIT = 30  # seconds for a bot check to clear itself and reload the page
    GIVE_UP_AFTER = 5  # pages in a row that never got past the bot check

    def __init__(self, workers: int, rate: float, browser: str = "chrome",
                 headless: bool = True, timeout: float = 60):
        self.workers = workers
        self.browser = browser
        self.headless = headless
        self.timeout_ms = timeout * 1000
        self.limiter = RateLimiter(rate)
        self.jobs: queue.Queue = queue.Queue()
        self.threads: list[threading.Thread] = []
        self.lock = threading.Lock()
        self.blocked_streak = 0
        self.given_up = False

    def start(self) -> None:
        try:
            import playwright.sync_api  # noqa: F401
        except ImportError:
            raise SetupError("Playwright is not installed: pip install playwright") from None
        for i in range(self.workers):
            ready = threading.Event()
            failure: list[Exception] = []
            thread = threading.Thread(target=self._work, args=(ready, failure),
                                      name=f"browser-{i}", daemon=True)
            thread.start()
            ready.wait()
            if failure:
                self.close()
                raise SetupError(f"could not start the browser ({self.browser}): {failure[0]}")
            self.threads.append(thread)

    def get(self, url: str) -> bytes:
        if self.given_up:
            raise Blocked("skipped: Zacks kept answering with a bot check")
        future: Future = Future()
        self.jobs.put((url, future))
        return future.result()

    def close(self) -> None:
        while True:  # cancel the pages that haven't started loading
            try:
                job = self.jobs.get_nowait()
            except queue.Empty:
                break
            if job is not None:
                job[1].cancel()
        for _ in self.threads:
            self.jobs.put(None)
        for thread in self.threads:
            thread.join(timeout=60)
        self.threads = []

    def _launch_options(self) -> dict:
        options: dict = {"headless": self.headless,
                         "args": ["--disable-blink-features=AutomationControlled"]}
        if self.browser == "chrome":
            options["channel"] = "chrome"  # the installed Google Chrome
        elif self.browser != "chromium":  # "chromium" is Playwright's own download
            options["executable_path"] = self.browser
        return options

    def _work(self, ready: threading.Event, failure: list[Exception]) -> None:
        from playwright.sync_api import sync_playwright
        try:
            with sync_playwright() as playwright:
                try:
                    browser = playwright.chromium.launch(**self._launch_options())
                    page = browser.new_page()
                except Exception as exc:
                    failure.append(exc)
                    return
                finally:
                    ready.set()
                self._serve(browser, page)
                browser.close()
        except Exception as exc:
            if not ready.is_set():
                failure.append(exc)
                ready.set()
            log.debug("Browser worker stopped", exc_info=True)

    def _serve(self, browser, page) -> None:
        while (job := self.jobs.get()) is not None:
            url, future = job
            if not future.set_running_or_notify_cancel():
                continue
            try:
                if self.given_up:
                    raise Blocked("skipped: Zacks kept answering with a bot check")
                if page.is_closed():
                    page = browser.new_page()
                future.set_result(self._load(page, url))
            except Exception as exc:
                future.set_exception(exc)

    def _load(self, page, url: str) -> bytes:
        self.limiter.wait()
        try:
            page.goto(url, wait_until="domcontentloaded", timeout=self.timeout_ms)
        except Exception:
            if page.url in ("", "about:blank"):
                raise  # nothing loaded at all
            # Otherwise the bot check's own reload interrupted it; keep watching the page.
        deadline = time.monotonic() + self.BOT_CHECK_WAIT
        while True:
            try:
                html = page.content()
            except Exception:  # the page is reloading
                html = ""
            if html and not _BOT_CHECK.search(html):
                self._record(blocked=False)
                return html.encode("utf-8")
            if time.monotonic() > deadline:
                self._record(blocked=True)
                raise Blocked("answered with a bot check that did not clear")
            page.wait_for_timeout(500)

    def _record(self, blocked: bool) -> None:
        with self.lock:
            self.blocked_streak = self.blocked_streak + 1 if blocked else 0
            if self.blocked_streak >= self.GIVE_UP_AFTER and not self.given_up:
                self.given_up = True
                log.error("Zacks keeps answering with a bot check, even in a browser; skipping "
                          "it for the rest of this run. Run again later to fill in the gaps.")


class Fetcher:
    """A cache in front of an HttpClient or BrowserPool, so reruns and resumed runs are fast."""

    def __init__(self, http: HttpClient | BrowserPool, cache: DiskCache, refresh: bool = False):
        self.http = http
        self.cache = cache
        self.refresh = refresh

    def get(self, url: str, max_age: float = HOLDINGS_TTL) -> bytes:
        return self._cached(url, max_age, lambda: self.http.get(url))

    def post_json(self, url: str, payload: dict, headers: dict | None = None,
                  max_age: float = HOLDINGS_TTL) -> bytes:
        return self._cached(request_key(url, payload), max_age,
                            lambda: self.http.post_json(url, payload, headers))

    def _cached(self, key: str, max_age: float, download: Callable[[], bytes]) -> bytes:
        if not self.refresh:
            cached = self.cache.read(key, max_age)
            if cached is not None:
                return cached
        data = download()
        self.cache.write(key, data)
        return data

    def forget(self, key: str) -> None:
        """Drop a cached response, e.g. an error page that should not be served again."""
        self.cache.delete(key)


def request_key(url: str, payload: dict | None = None) -> str:
    """What a response is cached under: the URL, plus a digest of what was posted to it."""
    if payload is None:
        return url
    body = json.dumps(payload, sort_keys=True).encode()
    return f"{url}#{hashlib.sha1(body).hexdigest()}"


# ----------------------------------------------------------------------- Universe


@dataclass(frozen=True)
class Etf:
    ticker: str
    name: str
    exchange: str


def parse_nasdaq_directory(text: str) -> list[Etf]:
    """ETFs from nasdaqtraded.txt (pipe-delimited, one row per US-listed security)."""
    etfs = []
    for row in csv.DictReader(io.StringIO(text), delimiter="|"):
        if row.get("ETF") != "Y" or row.get("Test Issue") == "Y":
            continue
        ticker = (row.get("Symbol") or "").strip()
        if not ticker:
            continue
        code = (row.get("Listing Exchange") or "").strip()
        name = (row.get("Security Name") or "").strip()
        etfs.append(Etf(ticker, name, EXCHANGES.get(code, code)))
    return sorted(etfs, key=lambda e: e.ticker)


def load_universe(fetcher: Fetcher) -> list[Etf]:
    raw = fetcher.get(NASDAQ_SYMBOL_DIRECTORY, max_age=REFERENCE_TTL)
    etfs = parse_nasdaq_directory(raw.decode("utf-8", errors="replace"))
    if not etfs:
        raise RuntimeError(f"No ETFs found in {NASDAQ_SYMBOL_DIRECTORY}; has its format changed?")
    return etfs


# ------------------------------------------------------------- Holdings tables

_MISSING = {"", "-", "--", "na", "n/a", "nan", "none", "null"}


def clean_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    text = html.unescape(str(value)).replace("\xa0", " ").strip()
    return "" if text.lower() in _MISSING else text


def parse_number(value: object) -> float | str:
    """'1,234.5', '$1,234.50', '2.66%', '(12)' -> float; blanks and dashes -> ''."""
    if isinstance(value, bool) or value is None:
        return ""
    if isinstance(value, (int, float)):
        return float(value)
    text = clean_text(value).replace(",", "").replace("$", "").replace("%", "").strip()
    negative = text.startswith("(") and text.endswith(")")
    try:
        number = float(text.strip("()"))
    except ValueError:
        return ""
    return -number if negative else number


_DATE_FORMATS = [
    (re.compile(r"\b\d{4}-\d{2}-\d{2}\b"), ("%Y-%m-%d",)),
    (re.compile(r"\b\d{1,2}/\d{1,2}/\d{4}\b"), ("%m/%d/%Y",)),
    (re.compile(r"\b\d{1,2}-[A-Za-z]{3}-\d{4}\b"), ("%d-%b-%Y",)),
    (re.compile(r"\b[A-Za-z]{3,9}\.? \d{1,2}, \d{4}\b"), ("%b %d, %Y", "%B %d, %Y")),
    (re.compile(r"\b\d{1,2} [A-Za-z]{3,9} \d{4}\b"), ("%d %b %Y", "%d %B %Y")),
    (re.compile(r"\b\d{1,2}/\d{1,2}/\d{2}\b"), ("%m/%d/%y",)),
]


def parse_date(text: str) -> str:
    """First recognizable date in text, as YYYY-MM-DD ('' if none)."""
    for pattern, formats in _DATE_FORMATS:
        for match in pattern.finditer(text):
            candidate = match.group(0).replace(".", "")
            for fmt in formats:
                try:
                    return dt.datetime.strptime(candidate, fmt).date().isoformat()
                except ValueError:
                    pass
    return ""


def parse_maturity(text: str) -> str:
    """A maturity date as YYYY-MM-DD. Pools of mortgages give a range ('11/01/2027-10/01/2056'),
    which becomes an ISO 8601 interval: '2027-11-01/2056-10-01'."""
    parts = re.split(r"(?<=\d)\s*-\s*(?=\d{1,2}/\d{1,2}/\d{4})", text)
    dates = [parse_date(part) for part in parts]
    return "/".join(dates) if all(dates) else text


def find_as_of(text: str) -> str:
    """The date following 'as of' in text, as YYYY-MM-DD ('' if none)."""
    for match in re.finditer(r"as of", text, re.IGNORECASE):
        found = parse_date(text[match.end():match.end() + 40])
        if found:
            return found
    return ""


# Issuer column headers, normalized by _header_key, mapped to output fields.
COLUMN_ALIASES = {
    "holding_name": ["name", "securityname", "security", "description", "securitydescription",
                     "holding", "holdingname", "issuername", "issuer", "company"],
    "ticker": ["ticker", "symbol", "tickersymbol", "stockticker"],
    "cusip": ["cusip"],
    "isin": ["isin"],
    "sedol": ["sedol", "sedolnumber"],
    "other_id": ["identifier", "securityid"],
    # Not "type": iShares files with both put e.g. "EQUITY" in Type, "Equity" in Asset Class.
    "asset_class": ["assetclass", "securitytype", "assettype"],
    "sector": ["sector", "gicssector", "industry", "classification"],
    "country": ["location", "country", "countryofrisk", "locationofrisk"],
    "currency": ["currency", "localcurrency", "marketcurrency", "basecurrency"],
    "shares": ["shares", "sharesheld", "quantity", "sharesquantity", "parvalue", "par", "units",
               "sharespar", "nominal", "principal", "sharescontracts"],
    "price": ["price", "marketprice", "closingprice"],
    # First Trust's "Market Value / Notional Value" holds notional amounts for derivatives;
    # add_value_weights() moves those to notional_value.
    "market_value": ["marketvalue", "marketvalueusd", "marketvalueus", "basemarketvalue", "value",
                     "marketvaluenotionalvalue"],
    # ProShares: "Exposure Value (Notional + GL)"
    "notional_value": ["notionalvalue", "exposurevaluenotionalgl"],
    "weight_pct": ["weightpct", "weight", "weighting", "weightings", "marketweight",
                   "exposureweight", "pctofnetassets", "percentofnetassets",
                   "pctnetassets", "netassetspct", "portfolioweight", "pctoffund",
                   "holdingspercent"],
    "maturity_date": ["maturity", "maturitydate"],
    "coupon_rate": ["couponpct", "coupon", "couponrate"],
    # Not an output column: files that repeat the holdings date on every row (see rows_to_holdings).
    "as_of_date": ["date", "asofdate", "tradedate"],
}
_ALIAS_TO_FIELD = {alias: f for f, aliases in COLUMN_ALIASES.items() for alias in aliases}


def _header_key(cell: object) -> str:
    """'Weight (%)' -> 'weightpct', 'Market Value ($)' -> 'marketvalue'."""
    return re.sub(r"[^a-z0-9]", "", clean_text(cell).lower().replace("%", "pct"))


def _map_header(cells: list[str], overrides: dict[str, str]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    for i, cell in enumerate(cells):
        key = _header_key(cell)
        target = overrides.get(key) or _ALIAS_TO_FIELD.get(key)
        if target and target not in mapping.values():
            mapping[i] = target
    return mapping


def _is_header(mapping: dict[int, str]) -> bool:
    fields = set(mapping.values())
    return bool(fields & {"holding_name", "ticker"}) and bool(
        fields & {"weight_pct", "shares", "market_value"}
    )


def normalize_holding(raw: dict) -> dict:
    holding = {}
    for name in HOLDING_FIELDS:
        value = raw.get(name)
        holding[name] = parse_number(value) if name in NUMERIC_FIELDS else clean_text(value)
    holding["maturity_date"] = parse_maturity(holding["maturity_date"])
    return holding


def rows_to_holdings(
    rows: Iterable[Sequence[object]], overrides: dict[str, str] | None = None
) -> tuple[list[dict], str]:
    """Holdings from a spreadsheet-like table, plus the text found above its header.

    Issuer files put a few lines of fund details (name, as-of date) above the
    table and disclaimers below it. The header row is the first row whose
    columns name a security and a weight, quantity or value; the table ends at
    the first row with fewer than two filled cells. ``overrides`` remaps
    ambiguous headers, e.g. SSGA's "Identifier" column holds CUSIPs.

    Some files (ARK, Pacer, Innovator) have no preamble but a date column; its
    first value is added to the preamble as "As of <date>".
    """
    overrides = overrides or {}
    preamble: list[str] = []
    mapping: dict[int, str] | None = None
    holdings = []
    row_date = ""
    for row in rows:
        cells = [clean_text(c) for c in row]
        if mapping is None:
            candidate = _map_header(cells, overrides)
            if _is_header(candidate):
                mapping = candidate
            else:
                preamble.append(" ".join(c for c in cells if c))
            continue
        if sum(1 for c in cells if c) < 2:
            break
        raw = {name: cells[i] for i, name in mapping.items() if i < len(cells)}
        row_date = row_date or raw.get("as_of_date", "")
        holdings.append(normalize_holding(raw))
    if mapping is None:
        raise NotAvailable("no holdings table found")
    if row_date:
        preamble.insert(0, f"As of {row_date}")
    return holdings, "\n".join(preamble)


def html_table_rows(table) -> list[list[str]]:
    """The text of each cell of a BeautifulSoup table, row by row."""
    return [[cell.get_text(" ", strip=True) for cell in tr.find_all(["td", "th"])]
            for tr in table.find_all("tr")]


def add_value_weights(rows: list[dict]) -> str:
    """Fill value_weight_pct: each holding's market value as a percentage of the fund's.

    Issuers round published weights to 0.01%, which hides most of a fund with
    thousands of small positions: 10,837 of AGG's 13,418 bonds show 0.00%, together
    21.6% of the fund. Short positions and written options have negative market
    values, so their weights are negative and the total is still 100% of net assets.

    Derivatives don't always fit: First Trust puts a future's notional amount in its
    market value column ("Market Value / Notional Value"), and American Century
    weights a future by its exposure while its market value is the day's gain or
    loss. So each row is checked against what its published weight allows, given the
    fund's net assets implied by the larger positions (leaving out those far from the
    median, such as derivatives). A market value too large for its weight is a notional
    amount and moves to notional_value; one too small is kept. The remaining market
    values must add up to those net assets. Holdings with neither a market value nor a weight
    (worthless warrants, say) are left out. Returns why value weights were not
    computed, or "" (also when a source has no market values).
    """
    valued = [r for r in rows if isinstance(r["market_value"], float)]
    if not valued:
        return ""
    if any(isinstance(r["weight_pct"], float) and r["weight_pct"] != 0
           for r in rows if not isinstance(r["market_value"], float)):
        return "value weights not computed: some holdings have no market value"
    weighted = [r for r in valued if isinstance(r["weight_pct"], float) and r["weight_pct"] != 0]
    # Larger positions lose least to rounding; equal-weight funds have none, so use them all.
    sized = [r for r in weighted if abs(r["weight_pct"]) >= 0.25]
    if sum(r["weight_pct"] for r in sized) < 10:
        sized = weighted
    if sum(r["weight_pct"] for r in sized) < 10:
        return "value weights not computed: too few published weights to check them against"
    implied = {id(r): r["market_value"] / r["weight_pct"] * 100 for r in sized}
    typical = statistics.median_low(implied.values())
    # Derivatives weighted by exposure, or valued at their notional amount, are far off.
    fitting = [r for r in sized if abs(implied[id(r)] / typical - 1) <= 0.1]
    net_assets = (sum(r["market_value"] for r in fitting)
                  / sum(r["weight_pct"] for r in fitting) * 100)
    notional = {id(r) for r in valued if _holds_notional_amount(r, net_assets)}
    total = sum(r["market_value"] for r in valued if id(r) not in notional)
    if abs(total / net_assets - 1) > 0.01:
        return (f"value weights not computed: market values add up to {total / net_assets:.1%} "
                "of the net assets implied by the published weights")
    for row in valued:
        if id(row) in notional:
            row["notional_value"] = row["notional_value"] or row["market_value"]
            row["market_value"] = ""
        else:
            row["value_weight_pct"] = round(row["market_value"] / total * 100, 6)
    return ""


def _holds_notional_amount(row: dict, net_assets: float) -> bool:
    """Is the market value far larger than the published weight allows?"""
    if not isinstance(row["weight_pct"], float):
        return False
    computed = row["market_value"] / net_assets * 100
    # Published weights are rounded to 0.01%; allow for that and for the estimate of net assets.
    return abs(computed) - abs(row["weight_pct"]) > max(0.01, 0.02 * abs(row["weight_pct"]))


def read_csv_rows(data: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))))


def read_xlsx_rows(data: bytes) -> list[list[object]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
    finally:
        workbook.close()


# ---------------------------------------------------------------------- Sources


@dataclass
class Holdings:
    source: str
    url: str
    as_of: str
    rows: list[dict]


class Source:
    name = ""

    def __init__(self, fetcher: Fetcher):
        self.fetcher = fetcher

    def prepare(self) -> None:
        """One-time setup before scraping starts, e.g. loading a fund list."""

    def handles(self, etf: Etf) -> bool:
        raise NotImplementedError

    def fetch(self, etf: Etf) -> Holdings:
        raise NotImplementedError

    def download(self, url: str, parse: Callable[[bytes], T], payload: dict | None = None,
                 headers: dict | None = None) -> T:
        """Fetch url (or post payload to it as JSON) and parse the response. A response
        that can't be parsed is usually an error or bot-check page: it is dropped from
        the cache, and the error says what it was."""
        if payload is None:
            data = self.fetcher.get(url)
        else:
            data = self.fetcher.post_json(url, payload, headers)
        try:
            return parse(data)
        except Exception as exc:
            self.fetcher.forget(request_key(url, payload))
            reason = str(exc) if isinstance(exc, NotAvailable) else f"{type(exc).__name__}: {exc}"
            raise NotAvailable(f"{reason}; {describe_response(data)}") from None


def describe_response(data: bytes) -> str:
    """'got an HTML page titled ...' or the first line of an unexpected response."""
    text = data[:50_000].decode("utf-8", errors="replace")
    title = re.search(r"<title[^>]*>(.*?)</title>", text, re.IGNORECASE | re.DOTALL)
    if title:
        return f"got a {len(data):,}-byte HTML page titled {clean_text(title.group(1))[:80]!r}"
    first = next((line.strip() for line in text.splitlines() if line.strip()), "")
    return f"got {len(data):,} bytes starting {first[:80]!r}"


class IsharesSource(Source):
    name = "ishares"
    FUND_LIST = (
        "https://www.ishares.com/us/product-screener/product-screener-v3.1.jsn?"
        "dcrPath=/templatedata/config/product-screener-v3/data/en/us-ishares/"
        "ishares-product-screener-backend-config&siteEntryPassthrough=true"
    )
    HOLDINGS = "https://www.ishares.com{page}/latest-holdings.csv"

    def __init__(self, fetcher: Fetcher):
        super().__init__(fetcher)
        self.pages: dict[str, str] = {}

    def prepare(self) -> None:
        try:
            funds = json.loads(self.fetcher.get(self.FUND_LIST, max_age=REFERENCE_TTL))
        except Exception as exc:
            log.warning("iShares fund list unavailable (%s); iShares ETFs will use the "
                        "fallback source", exc)
            return
        self.pages = parse_ishares_fund_list(funds)
        log.info("iShares: %d funds", len(self.pages))

    def handles(self, etf: Etf) -> bool:
        return etf.ticker in self.pages

    def fetch(self, etf: Etf) -> Holdings:
        url = self.HOLDINGS.format(page=self.pages[etf.ticker])
        holdings, preamble = self.download(url, lambda d: rows_to_holdings(read_csv_rows(d)))
        return Holdings(self.name, url, find_as_of(preamble), holdings)


def parse_ishares_fund_list(funds: dict) -> dict[str, str]:
    """{product id: {...}} from the iShares product screener -> {ticker: product page path}."""
    pages = {}
    for fund in funds.values():
        ticker = (fund.get("localExchangeTicker") or "").strip().upper()
        page = (fund.get("productPageUrl") or "").strip()
        if ticker and page.startswith("/"):
            pages[ticker] = page.rstrip("/")
    return pages


class SsgaSource(Source):
    name = "ssga"
    HOLDINGS = ("https://www.ssga.com/library-content/products/fund-data/etfs/us/"
                "holdings-daily-us-en-{ticker}.xlsx")
    NAME = re.compile(r"\bSPDR\b|^State Street\b", re.IGNORECASE)

    def handles(self, etf: Etf) -> bool:
        return bool(self.NAME.search(etf.name))

    def fetch(self, etf: Etf) -> Holdings:
        url = self.HOLDINGS.format(ticker=etf.ticker.lower())
        holdings, preamble = self.download(
            url, lambda d: rows_to_holdings(read_xlsx_rows(d), {"identifier": "cusip"}))
        return Holdings(self.name, url, find_as_of(preamble), holdings)


class InvescoSource(Source):
    name = "invesco"
    # Most funds are keyed by CUSIP (only a few, like QQQ, by ticker); the fund page says which.
    FUND_PAGE = ("https://www.invesco.com/us/financial-products/etfs/product-detail"
                 "?audienceType=Investor&ticker={ticker}")
    HOLDINGS = ("https://dng-api.invesco.com/cache/v1/accounts/en_US/shareclasses/{fund_id}"
                "/holdings/fund?idType={id_type}&interval=monthly&productType=ETF")

    def handles(self, etf: Etf) -> bool:
        return etf.name.lower().startswith("invesco")

    def fetch(self, etf: Etf) -> Holdings:
        page_url = self.FUND_PAGE.format(ticker=etf.ticker)
        id_type, fund_id = self.download(
            page_url, lambda d: parse_invesco_fund_id(d.decode("utf-8", errors="replace")))
        url = self.HOLDINGS.format(fund_id=fund_id, id_type=id_type)
        as_of, holdings = self.download(url, lambda d: parse_invesco(json.loads(d)))
        return Holdings(self.name, url, as_of, holdings)


def parse_invesco_fund_id(page: str) -> tuple[str, str]:
    """(id type, id) from a fund page's meta tags, e.g. ("cusip", "46137V357")."""
    meta = dict(re.findall(r'<meta\s+name="([^"]+)"\s+content="([^"]*)"', page))
    id_type = meta.get("uniqueIdentifier", "")
    fund_id = meta.get(id_type, "")
    if not id_type or not fund_id:
        raise NotAvailable("fund page has no fund identifier")
    return id_type, fund_id


def parse_invesco(data: dict) -> tuple[str, list[dict]]:
    holdings = [
        normalize_holding({
            "holding_name": h.get("issuerName") or h.get("securityName"),
            "ticker": h.get("ticker"),
            "cusip": h.get("cusip"),
            "isin": h.get("isin"),
            "sedol": h.get("sedol"),
            "asset_class": h.get("securityTypeName"),
            "sector": h.get("sectorName"),
            "country": h.get("countryName"),
            "currency": h.get("currency"),
            "shares": h.get("units"),
            "price": h.get("price"),
            "market_value": h.get("marketValue", h.get("marketValueBase")),
            "weight_pct": h.get("percentageOfTotalNetAssets"),
            "maturity_date": h.get("maturityDate"),
            "coupon_rate": h.get("couponRate", h.get("coupon")),
        })
        for h in data.get("holdings") or []
    ]
    as_of = parse_date(str(data.get("effectiveBusinessDate") or data.get("effectiveDate") or ""))
    return as_of, holdings


class FirstTrustSource(Source):
    name = "first_trust"
    HOLDINGS = "https://www.ftportfolios.com/Retail/Etf/EtfHoldings.aspx?Ticker={ticker}"
    NAME = re.compile(r"^(First Trust|FT Vest)\b", re.IGNORECASE)

    def handles(self, etf: Etf) -> bool:
        return bool(self.NAME.search(etf.name))

    def fetch(self, etf: Etf) -> Holdings:
        url = self.HOLDINGS.format(ticker=etf.ticker)
        as_of, holdings = self.download(
            url, lambda d: parse_first_trust(d.decode("utf-8", errors="replace")))
        return Holdings(self.name, url, as_of, holdings)


def parse_first_trust(page: str) -> tuple[str, list[dict]]:
    soup = BeautifulSoup(page, "html.parser")
    title = soup.find(id=re.compile(r"lblHoldingsTitle$"))
    as_of = find_as_of(title.get_text(" ", strip=True)) if title else ""
    holdings = []
    for table in soup.select("table.fundSilverGrid"):
        try:
            # Here "Identifier" is the exchange ticker; CUSIPs have their own column.
            table_holdings, _ = rows_to_holdings(html_table_rows(table), {"identifier": "ticker"})
        except NotAvailable:
            continue  # e.g. a futures exposure summary table
        holdings.extend(table_holdings)
    if not holdings:
        raise NotAvailable("no holdings table found")
    return as_of, holdings


class DatedFileSource(Source):
    """Holdings files named by date ({date} in HOLDINGS): the most recent one is used."""

    HOLDINGS = ""
    LOOKBACK_BUSINESS_DAYS = 5

    def __init__(self, fetcher: Fetcher, today: dt.date | None = None):
        super().__init__(fetcher)
        self.today = today

    def file_ticker(self, etf: Etf) -> str:
        return etf.ticker

    def parse(self, data: bytes) -> tuple[list[dict], str]:
        return rows_to_holdings(read_csv_rows(data))

    def fetch(self, etf: Etf) -> Holdings:
        for day in recent_business_days(self.today or dt.date.today(), self.LOOKBACK_BUSINESS_DAYS):
            url = self.HOLDINGS.format(ticker=self.file_ticker(etf), date=day)
            try:
                holdings, preamble = self.download(url, self.parse)
            except NotFound:
                continue
            except requests.HTTPError as exc:
                # Some file stores answer 403 for a file that doesn't exist (yet).
                if exc.response is not None and exc.response.status_code == 403:
                    continue
                raise
            return Holdings(self.name, url, find_as_of(preamble) or day.isoformat(), holdings)
        raise NotAvailable(
            f"no holdings file in the last {self.LOOKBACK_BUSINESS_DAYS} business days")


class GlobalXSource(DatedFileSource):
    name = "global_x"
    HOLDINGS = "https://assets.globalxetfs.com/funds/holdings/{ticker}_full-holdings_{date:%Y%m%d}.csv"

    def handles(self, etf: Etf) -> bool:
        return etf.name.lower().startswith("global x")

    def file_ticker(self, etf: Etf) -> str:
        return etf.ticker.lower()


class DimensionalSource(DatedFileSource):
    name = "dimensional"
    HOLDINGS = "https://tools-blob.dimensional.com/etf/{date:%Y%m%d}/{ticker}.csv"

    def handles(self, etf: Etf) -> bool:
        return etf.name.lower().startswith("dimensional")

    def parse(self, data: bytes) -> tuple[list[dict], str]:
        # "identifier" holds CUSIPs, and weights are fractions: 0.0567 means 5.67%.
        holdings, preamble = rows_to_holdings(read_csv_rows(data), {"identifier": "cusip"})
        for holding in holdings:
            if isinstance(holding["weight_pct"], float):
                holding["weight_pct"] = round(holding["weight_pct"] * 100, 10)
        return holdings, preamble


def recent_business_days(today: dt.date, count: int) -> list[dt.date]:
    days, day = [], today
    while len(days) < count:
        if day.weekday() < 5:
            days.append(day)
        day -= dt.timedelta(days=1)
    return days


class CsvSource(Source):
    """Holdings in a CSV file at a URL made from the ticker; funds are matched by name."""

    HOLDINGS = ""
    NAME: re.Pattern

    def handles(self, etf: Etf) -> bool:
        return bool(self.NAME.search(etf.name))

    def fetch(self, etf: Etf) -> Holdings:
        url = self.HOLDINGS.format(ticker=etf.ticker)
        holdings, preamble = self.download(url, lambda d: rows_to_holdings(read_csv_rows(d)))
        return Holdings(self.name, url, find_as_of(preamble), holdings)


class PacerSource(CsvSource):
    name = "pacer"
    HOLDINGS = "https://www.paceretfs.com/products/holdings_download/{ticker}"
    NAME = re.compile(r"^Pacer\b", re.IGNORECASE)


class DirexionSource(CsvSource):
    name = "direxion"
    HOLDINGS = "https://www.direxion.com/holdings/{ticker}.csv"
    NAME = re.compile(r"^Direxion\b", re.IGNORECASE)


class InnovatorSource(Source):
    """Innovator publishes one file with the holdings of all its funds."""

    name = "innovator"
    HOLDINGS = "https://www.innovatoretfs.com/etf/xt_holdings.csv"

    def __init__(self, fetcher: Fetcher):
        super().__init__(fetcher)
        self.lock = threading.Lock()
        self.funds: dict[str, list[list[str]]] | None = None
        self.error = ""

    def handles(self, etf: Etf) -> bool:
        return etf.name.lower().startswith("innovator")

    def fetch(self, etf: Etf) -> Holdings:
        rows = self._funds().get(etf.ticker)
        if not rows:
            raise NotAvailable("not in the holdings file")
        holdings, preamble = rows_to_holdings(rows)
        return Holdings(self.name, self.HOLDINGS, find_as_of(preamble), holdings)

    def _funds(self) -> dict[str, list[list[str]]]:
        """The file split by fund, downloaded once by whichever thread asks first."""
        with self.lock:
            if self.funds is None and not self.error:
                try:
                    self.funds = self.download(
                        self.HOLDINGS, lambda d: split_by_fund(read_csv_rows(d), "account"))
                except NotAvailable as exc:
                    self.error = str(exc)
                except Exception as exc:
                    self.error = f"{type(exc).__name__}: {exc}"
            if self.error:
                raise NotAvailable(self.error)
            return self.funds


def split_by_fund(rows: list[list[str]], fund_column: str) -> dict[str, list[list[str]]]:
    """{fund ticker: [header, *its rows]} from a file listing many funds' holdings."""
    for start, row in enumerate(rows):
        keys = [_header_key(cell) for cell in row]
        if fund_column in keys:
            column = keys.index(fund_column)
            break
    else:
        raise NotAvailable(f"no {fund_column!r} column")
    funds: dict[str, list[list[str]]] = {}
    for row in rows[start + 1:]:
        fund = row[column].strip().upper() if column < len(row) else ""
        if fund:
            funds.setdefault(fund, [rows[start]]).append(row)
    return funds


class VanguardSource(Source):
    """Vanguard publishes holdings as of each month end, a few weeks after it."""

    name = "vanguard"
    HOLDINGS = ("https://investor.vanguard.com/vmf/api/{ticker}/portfolio-holding/{kind}"
                "?start={start}&count={count}")
    KINDS = ("stock", "bond")  # listed separately; a fund has none of one kind or the other
    PAGE_SIZE = 3000

    def handles(self, etf: Etf) -> bool:
        return etf.name.lower().startswith("vanguard")

    def fetch(self, etf: Etf) -> Holdings:
        source_url, as_of, holdings = "", "", []
        for kind in self.KINDS:
            start, total = 1, 1
            while start <= total:
                url = self.HOLDINGS.format(ticker=etf.ticker, kind=kind, start=start,
                                           count=self.PAGE_SIZE)
                total, page_as_of, rows = self.download(
                    url, lambda d: parse_vanguard(json.loads(d), kind))
                if not rows:
                    break
                source_url, as_of = source_url or url, as_of or page_as_of
                holdings.extend(rows)
                start += len(rows)
        return Holdings(self.name, source_url, as_of, holdings)


def parse_vanguard(data: dict, kind: str) -> tuple[int, str, list[dict]]:
    """(holdings in total, as-of date, holdings on this page) from Vanguard's holdings API."""
    holdings = []
    for h in (data.get("fund") or {}).get("entity") or []:
        holdings.append(normalize_holding({
            "holding_name": h.get("longName") or h.get("shortName"),
            "ticker": h.get("ticker"),
            "cusip": h.get("cusip"),
            "isin": h.get("isin"),
            "sedol": h.get("sedol"),
            "asset_class": kind.capitalize(),
            # Bonds list their par value as faceAmount; sharesHeld is 0.
            "shares": h.get("faceAmount") or h.get("sharesHeld"),
            "market_value": h.get("marketValue"),
            "notional_value": h.get("notionalValue") if parse_number(h.get("notionalValue"))
            else "",
            "weight_pct": h.get("percentWeight"),
            "maturity_date": h.get("maturityDate"),
            "coupon_rate": h.get("couponRate"),
        }))
    as_of = parse_date(str(data.get("asOfDate") or "").split("T")[0])  # 2026-08-31T00:00:00-04:00
    return int(data.get("size") or 0), as_of, holdings


class ProSharesSource(Source):
    name = "proshares"
    # Fund pages list every holding. Leveraged and inverse funds have their own path.
    PAGES = ("https://www.proshares.com/our-etfs/leveraged-and-inverse/{ticker}",
             "https://www.proshares.com/our-etfs/strategic/{ticker}")

    def handles(self, etf: Etf) -> bool:
        return etf.name.lower().startswith("proshares")

    def fetch(self, etf: Etf) -> Holdings:
        for page in self.PAGES:
            url = page.format(ticker=etf.ticker.lower())
            try:
                as_of, holdings = self.download(
                    url, lambda d: parse_proshares(d.decode("utf-8", errors="replace")))
            except NotFound:
                continue
            return Holdings(self.name, url, as_of, holdings)
        raise NotAvailable("no fund page")


def parse_proshares(page: str) -> tuple[str, list[dict]]:
    """Holdings from a ProShares fund page. Swaps and futures have an exposure value
    (notional_value) and a weight but no market value; their gains and losses are in
    the "Net Other Assets" row."""
    soup = BeautifulSoup(page, "html.parser")
    table = soup.find("table", id="holdings")
    if table is None:
        raise NotAvailable("no holdings table found")
    section = table.find_parent("section") or table
    holdings, _ = rows_to_holdings(html_table_rows(table))
    return find_as_of(section.get_text(" ", strip=True)), holdings


class SchwabSource(Source):
    name = "schwab"
    HOLDINGS = "https://www.schwabassetmanagement.com/allholdings/{ticker}?page={page}"

    def handles(self, etf: Etf) -> bool:
        return etf.name.lower().startswith("schwab")

    def fetch(self, etf: Etf) -> Holdings:
        as_of, holdings, page, total = "", [], 0, 1
        while len(holdings) < total:  # 100 holdings per page
            url = self.HOLDINGS.format(ticker=etf.ticker, page=page)
            page_as_of, total, rows = self.download(
                url, lambda d: parse_schwab(d.decode("utf-8", errors="replace")))
            if not rows:
                break
            as_of = as_of or page_as_of
            holdings.extend(rows)
            page += 1
        return Holdings(self.name, self.HOLDINGS.format(ticker=etf.ticker, page=0), as_of, holdings)


def parse_schwab(page: str) -> tuple[str, int, list[dict]]:
    """(as-of date, holdings in total, holdings on this page) from a Schwab all-holdings page.
    Market values are left out: the page rounds them to a few digits ("$5.2B")."""
    soup = BeautifulSoup(page, "html.parser")
    table = soup.select_one("table.table-funds")
    if table is None:
        raise NotAvailable("no holdings table found")
    holdings, _ = rows_to_holdings(html_table_rows(table),
                                   {"fundname": "holding_name", "pctofassets": "weight_pct"})
    for holding in holdings:
        holding["market_value"] = ""
    as_of = soup.select_one(".view__header--as-of-date")
    shown = soup.select_one(".view__header--results")  # "Displaying 1 - 100 of 102"
    total = re.search(r"\bof\s+([\d,]+)", shown.get_text(" ", strip=True)) if shown else None
    return (find_as_of(as_of.get_text(" ", strip=True)) if as_of else "",
            int(total.group(1).replace(",", "")) if total else len(holdings),
            holdings)


class VanEckSource(Source):
    name = "vaneck"
    # The path names the fund's asset class. With the wrong one, the site redirects to
    # the fund's overview page, so bond funds try "income" first and the others "equity".
    HOLDINGS = "https://www.vaneck.com/us/en/etf/{category}/{ticker}/holdings/download/xlsx/"
    INCOME = re.compile(r"\b(Bond|Muni|Treasury|Yield|Income|Credit|Loan|CLO|Mortgage)\b",
                        re.IGNORECASE)

    def handles(self, etf: Etf) -> bool:
        return etf.name.lower().startswith("vaneck")

    def fetch(self, etf: Etf) -> Holdings:
        categories = ("income", "equity") if self.INCOME.search(etf.name) else ("equity", "income")
        failures = []
        for category in categories:
            url = self.HOLDINGS.format(category=category, ticker=etf.ticker.lower())
            try:
                holdings, preamble = self.download(url, lambda d: rows_to_holdings(
                    read_xlsx_rows(d), {"identifierfigi": "other_id"}))
            except NotAvailable as exc:
                failures.append(f"{category}: {exc}")
                continue
            # The workbook starts "Daily Holdings (%) 10/13/2026", without "as of".
            return Holdings(self.name, url, find_as_of(preamble) or parse_date(preamble), holdings)
        raise NotAvailable("; ".join(failures))


class ArkSource(Source):
    name = "ark"
    FUND_PAGE = "https://www.ark-funds.com/funds/{ticker}"
    NAME = re.compile(r"^ARK\b")  # not IGNORECASE: "Ark" could start other fund names

    def handles(self, etf: Etf) -> bool:
        return bool(self.NAME.search(etf.name))

    def fetch(self, etf: Etf) -> Holdings:
        # The holdings file is named after the fund ("ARK_INNOVATION_ETF_ARKK_HOLDINGS.csv"),
        # not always as the ETF list spells it, so take the link from the fund page.
        page_url = self.FUND_PAGE.format(ticker=etf.ticker.lower())
        url = self.download(page_url, lambda d: find_ark_csv(d.decode("utf-8", errors="replace")))
        holdings, preamble = self.download(url, lambda d: rows_to_holdings(read_csv_rows(d)))
        return Holdings(self.name, url, find_as_of(preamble), holdings)


_ARK_CSV = re.compile(
    r"https://assets\.ark-funds\.com/fund-documents/funds-etf-csv/[^\"'<>\s]+\.csv")


def find_ark_csv(page: str) -> str:
    match = _ARK_CSV.search(page)
    if not match:
        raise NotAvailable("fund page links no holdings file")
    return html.unescape(match.group(0))


class AmericanCenturySource(Source):
    """American Century's sites show each ETF's holdings on a page addressed by an internal
    fund number, which the fund list gives. Avantis is American Century's other brand."""

    name = "american_century"
    FUND_LIST = "https://www.americancentury.com/invest/funds/?type=ETF"
    HOLDINGS = "https://www.americancentury.com/invest/funds/total-holdings/{number}/?type=etf"

    def __init__(self, fetcher: Fetcher):
        super().__init__(fetcher)
        self.numbers: dict[str, str] = {}

    def prepare(self) -> None:
        try:
            page = self.fetcher.get(self.FUND_LIST, max_age=REFERENCE_TTL)
        except Exception as exc:
            log.warning("%s fund list unavailable (%s); its ETFs will use the fallback source",
                        self.name, exc)
            return
        self.numbers = parse_american_century_fund_list(page.decode("utf-8", errors="replace"))
        log.info("%s: %d ETFs", self.name, len(self.numbers))

    def handles(self, etf: Etf) -> bool:
        return etf.ticker in self.numbers

    def fetch(self, etf: Etf) -> Holdings:
        url = self.HOLDINGS.format(number=self.numbers[etf.ticker])
        as_of, holdings = self.download(
            url, lambda d: parse_american_century(d.decode("utf-8", errors="replace")))
        return Holdings(self.name, url, as_of, holdings)


class AvantisSource(AmericanCenturySource):
    name = "avantis"
    FUND_LIST = "https://www.avantisinvestors.com/avantis-investments/"
    HOLDINGS = ("https://www.avantisinvestors.com/avantis-investments/total-holdings/{number}/"
                "?type=etf")


# The pages are rendered from JavaScript object literals embedded in them, such as
# {ticker:"AVUS",productName:"Avantis U.S. Equity ETF",fundNumber:"114",productType:"ETF",...}
_JS_FIELD = re.compile(r'[{,](\w+):(?:"((?:[^"\\]|\\.)*)"|null)')


def _js_fields(text: str) -> dict[str, str]:
    """The string fields of the object literal text starts with (the first of each name)."""
    fields: dict[str, str] = {}
    for key, value in _JS_FIELD.findall(text):
        try:
            fields.setdefault(key, json.loads(f'"{value}"'))
        except ValueError:
            fields.setdefault(key, value)
    return fields


def parse_american_century_fund_list(page: str) -> dict[str, str]:
    """{ETF ticker: fund number} from an American Century or Avantis fund list page."""
    numbers = {}
    for match in re.finditer(r'\{ticker:"', page):
        fund = _js_fields(page[match.start():match.start() + 500])
        if fund.get("productType") == "ETF" and fund.get("fundNumber"):
            numbers[fund["ticker"].upper()] = fund["fundNumber"]
    return numbers


def js_array_items(text: str, start: int) -> list[str]:
    """The {...} items of the JavaScript array literal that starts at text[start] ("[")."""
    items, depth, item_start, in_string, i = [], 0, 0, False, start
    while i < len(text):
        char = text[i]
        if in_string:
            if char == "\\":
                i += 1
            elif char == '"':
                in_string = False
        elif char == '"':
            in_string = True
        elif char in "[{":
            depth += 1
            if char == "{" and depth == 2:
                item_start = i
        elif char in "]}":
            depth -= 1
            if char == "}" and depth == 1:
                items.append(text[item_start:i + 1])
            if depth == 0:
                return items
        i += 1
    raise NotAvailable("holdings list is cut off")


def parse_american_century(page: str) -> tuple[str, list[dict]]:
    """(as-of date, holdings) from an American Century or Avantis total-holdings page."""
    start = page.find("etfHoldings:[")
    if start < 0:
        raise NotAvailable("no holdings on page")
    holdings = []
    for item in js_array_items(page, start + len("etfHoldings:")):
        h = _js_fields(item)
        holdings.append(normalize_holding({
            "holding_name": h.get("name"),
            "ticker": h.get("ticker"),
            "cusip": h.get("cusip"),
            "isin": h.get("isin"),
            "sedol": h.get("sedol"),
            "asset_class": h.get("securityType"),
            "sector": h.get("sector"),
            "country": h.get("country"),
            "shares": h.get("shareQuantity"),
            "market_value": h.get("baseMarketValue"),
            "weight_pct": h.get("weight"),
            "maturity_date": h.get("maturityDate"),
            "coupon_rate": h.get("coupon"),
        }))
    as_of = re.search(r'etfHoldingsAsOfDate:"([^"]*)"', page)
    return parse_date(as_of.group(1)) if as_of else "", holdings


class FidelitySource(Source):
    """Fidelity's research site lists each ETF's holdings as of a month end, with weights
    only: the last month end for most funds, the one before for its active equity ETFs."""

    name = "fidelity"
    HOLDINGS = ("https://research2.fidelity.com/fidelity/screeners/etf/etfholdings.asp"
                "?symbol={ticker}&view=Holdings")

    def handles(self, etf: Etf) -> bool:
        return etf.name.lower().startswith("fidelity")

    def fetch(self, etf: Etf) -> Holdings:
        url = self.HOLDINGS.format(ticker=etf.ticker)
        as_of, holdings = self.download(
            url, lambda d: parse_fidelity(d.decode("utf-8", errors="replace")))
        return Holdings(self.name, url, as_of, holdings)


def parse_fidelity(page: str) -> tuple[str, list[dict]]:
    """(as-of date, holdings) from a Fidelity research holdings page. A fund it has no
    holdings for gets one row saying so, which ends the table: no holdings."""
    soup = BeautifulSoup(page, "html.parser")
    table = soup.select_one("table.results-table")
    if table is None:
        raise NotAvailable("no holdings table found")
    holdings, _ = rows_to_holdings(html_table_rows(table))
    header = soup.select_one(".num-results")  # "Basket Holdings: 298 AS OF 09/30/2026"
    return find_as_of(header.get_text(" ", strip=True)) if header else "", holdings


class GoldmanSachsSource(Source):
    """Goldman Sachs's fund pages load all holdings with a GraphQL query (a POST), which
    takes the fund's product number and share class ID (its CUSIP). Those come from the
    Fund Finder, which only fills in its list in the browser, so they are listed here:
    all 48 Goldman Sachs ETFs on 2026-10-21. A newer fund falls back to Zacks until it
    is added (see the README)."""

    name = "goldman_sachs"
    URL = "https://am.gs.com/services/funds"
    QUERY = """query getFundsDetail($fundDetailRequest: FundDetailRequest) {
  fundsDetail(fundDetailRequest: $fundDetailRequest) {
    ticker
    allHoldings {
      asAtDate
      holdings { holding ticker cusip isin sedol securityIdentifier numberOfShares
                 marketValue portfolio coupon maturityDate }
    }
  }
}"""
    FUNDS = {  # ticker: (product number, share class ID)
        "AAAU": ("PV103623", "38150K103"), "GBIL": ("PV102645", "381430529"),
        "GBND": ("PV109040", "38149W473"), "GCAL": ("PV105933", "38149W564"),
        "GCOR": ("PV103404", "38149W101"), "GCPB": ("PV110568", "38151N809"),
        "GDOC": ("PV103944", "38149W770"), "GEM": ("PV102391", "381430206"),
        "GEMD": ("PV102979", "381430388"), "GEMQ": ("PV110439", "38149W390"),
        "GGUS": ("PV105434", "38149W598"), "GHYB": ("PV102746", "381430453"),
        "GIEQ": ("PV110311", "38149W382"), "GIGB": ("PV102748", "381430479"),
        "GIGL": ("PV109041", "38149W465"), "GINC": ("PV110567", "38151N882"),
        "GIND": ("PV106306", "38149W481"), "GINN": ("PV103686", "38149W820"),
        "GMNY": ("PV105932", "38149W556"), "GMUB": ("PV105934", "38149W549"),
        "GPIQ": ("PV105259", "38149W630"), "GPIX": ("PV105258", "38149W622"),
        "GPRF": ("PV105611", "38149W127"), "GSC": ("PV105332", "38149W614"),
        "GSEE": ("PV103363", "381430164"), "GSEU": ("PV102392", "381430305"),
        "GSEW": ("PV102787", "381430438"), "GSGO": ("PV109510", "38149W440"),
        "GSID": ("PV103362", "381430180"), "GSIE": ("PV102390", "381430107"),
        "GSJY": ("PV102393", "381430404"), "GSLC": ("PV102394", "381430503"),
        "GSSC": ("PV102395", "381430602"), "GSST": ("PV103157", "381430230"),
        "GSUS": ("PV103365", "381430123"), "GSWO": ("PV104218", "38149W739"),
        "GTEK": ("PV103849", "38149W812"), "GTIP": ("PV102998", "381430362"),
        "GTOP": ("PV109507", "38149W432"), "GTPE": ("PV109407", "38149W457"),
        "GUMI": ("PV105931", "38149W572"), "GUSA": ("PV104276", "38150W107"),
        "GUSE": ("PV109508", "38149W424"), "GVIP": ("PV102634", "381430545"),
        "GVLE": ("PV109509", "38149W416"), "GVUS": ("PV105433", "38149W580"),
        "GXUS": ("PV104572", "38150W206"), "JUST": ("PV102957", "381430396"),
    }

    def handles(self, etf: Etf) -> bool:
        return etf.ticker in self.FUNDS

    def fetch(self, etf: Etf) -> Holdings:
        product, share_class = self.FUNDS[etf.ticker]
        payload = {
            "operationName": "getFundsDetail",
            "query": self.QUERY,
            "variables": {"fundDetailRequest": {
                "country": "us", "language": "en", "audience": "individual",
                "pvNumber": product, "shareClassId": share_class}},
        }
        as_of, holdings = self.download(
            self.URL, lambda d: parse_goldman_sachs(json.loads(d), etf.ticker),
            payload=payload, headers={"Origin": "https://am.gs.com"})
        return Holdings(self.name, self.URL, as_of, holdings)


def parse_goldman_sachs(data: dict, ticker: str) -> tuple[str, list[dict]]:
    """(as-of date, holdings) from Goldman Sachs's getFundsDetail response."""
    if data.get("errors"):
        messages = "; ".join(str(e.get("message", e)) for e in data["errors"])
        raise NotAvailable(f"the query was refused: {messages}")
    fund = (data.get("data") or {}).get("fundsDetail")
    if not fund:
        raise NotAvailable("no such fund")
    if fund.get("ticker") and fund["ticker"].upper() != ticker:
        raise NotAvailable(f"the IDs listed for {ticker} belong to {fund['ticker']}")
    all_holdings = fund.get("allHoldings") or {}
    holdings = [
        normalize_holding({
            "holding_name": h.get("holding"),
            "ticker": h.get("ticker"),
            "cusip": h.get("cusip"),
            "isin": h.get("isin"),
            "sedol": h.get("sedol"),
            "other_id": h.get("securityIdentifier"),
            "shares": h.get("numberOfShares"),
            "market_value": h.get("marketValue"),
            "weight_pct": h.get("portfolio"),
            "coupon_rate": h.get("coupon"),
            "maturity_date": h.get("maturityDate"),
        })
        for h in all_holdings.get("holdings") or []
    ]
    return parse_date(str(all_holdings.get("asAtDate") or "")), holdings


class ZacksSource(Source):
    name = "zacks"
    HOLDINGS = "https://www.zacks.com/funds/etf/{ticker}/holding"

    def handles(self, etf: Etf) -> bool:
        return True

    def fetch(self, etf: Etf) -> Holdings:
        url = self.HOLDINGS.format(ticker=etf.ticker)
        holdings = self.download(url, lambda d: parse_zacks(d.decode("utf-8", errors="replace")))
        return Holdings(self.name, url, "", holdings)


_ZACKS_DATA = re.compile(r"etf_holdings\.formatted_data\s*=\s*(\[.*?\])\s*;", re.DOTALL)
_TITLE_ATTR = re.compile(r'title="([^"]*)"')
_REL_ATTR = re.compile(r'rel="([^"]+)"')
_TAGS = re.compile(r"<[^>]+>")
_CUSIP = re.compile(r"^[0-9A-Z]{8}[0-9]$")


def parse_zacks(page: str) -> list[dict]:
    """Holdings from the JavaScript array embedded in a Zacks ETF holdings page.

    Each row is [name, symbol, shares, weight %, 52-week change, report link];
    long names are truncated in the text but complete in a title attribute.
    """
    match = _ZACKS_DATA.search(page)
    if not match:
        raise NotAvailable("no holdings on page")
    holdings = []
    for row in json.loads(match.group(1)):
        if len(row) < 4:
            continue
        name_cell, symbol_cell, shares, weight = (str(c) for c in row[:4])
        title = _TITLE_ATTR.search(name_cell)
        name = title.group(1) if title else _TAGS.sub("", name_cell)
        linked = _REL_ATTR.search(symbol_cell)
        symbol = clean_text(linked.group(1) if linked else _TAGS.sub("", symbol_cell))
        fields = {"holding_name": name, "shares": shares, "weight_pct": weight}
        if linked:
            fields["ticker"] = symbol
        elif _CUSIP.match(symbol):
            fields["cusip"] = symbol
        else:
            fields["other_id"] = symbol
        holdings.append(normalize_holding(fields))
    return holdings


ISSUER_SOURCES: dict[str, type[Source]] = {
    "ishares": IsharesSource,
    "ssga": SsgaSource,
    "invesco": InvescoSource,
    "first_trust": FirstTrustSource,
    "global_x": GlobalXSource,
    "vanguard": VanguardSource,
    "proshares": ProSharesSource,
    "innovator": InnovatorSource,
    "pacer": PacerSource,
    "schwab": SchwabSource,
    "vaneck": VanEckSource,
    "ark": ArkSource,
    "direxion": DirexionSource,
    "dimensional": DimensionalSource,
    "avantis": AvantisSource,
    "american_century": AmericanCenturySource,
    "fidelity": FidelitySource,
    "goldman_sachs": GoldmanSachsSource,
}


# ---------------------------------------------------------------------- Scraper


class HoldingsScraper:
    """Tries the ETF's issuer source first, then the fallback, and keeps the first success."""

    def __init__(self, issuers: list[Source], fallback: Source | None):
        self.issuers = issuers
        self.fallback = fallback

    def sources_for(self, etf: Etf) -> list[Source]:
        chain = [s for s in self.issuers if s.handles(etf)]
        if self.fallback is not None:
            chain.append(self.fallback)
        return chain

    def scrape(self, etf: Etf) -> tuple[dict, list[dict]]:
        summary = {col: "" for col in SUMMARY_COLUMNS}
        summary.update(ticker=etf.ticker, name=etf.name, exchange=etf.exchange)
        failures = []
        for source in self.sources_for(etf):
            try:
                result = source.fetch(etf)
                if not result.rows:
                    raise NotAvailable("no holdings listed")
            except NotAvailable as exc:
                failures.append(f"{source.name}: {exc}")
                continue
            except Exception as exc:  # one bad page must not stop the run
                log.debug("%s failed for %s", source.name, etf.ticker, exc_info=True)
                failures.append(f"{source.name}: {type(exc).__name__}: {exc}")
                continue

            value_weights_note = add_value_weights(result.rows)
            weights = [h["weight_pct"] for h in result.rows if isinstance(h["weight_pct"], float)]
            summary.update(
                status="ok",
                source=result.source,
                as_of_date=result.as_of,
                holdings_count=len(result.rows),
                weight_sum_pct=round(sum(weights), 2) if weights else "",
                source_url=result.url,
                detail="; ".join(failures + [value_weights_note] if value_weights_note
                                 else failures),
            )
            prefix = {"etf_ticker": etf.ticker, "etf_name": etf.name,
                      "source": result.source, "as_of_date": result.as_of}
            rows = [{**prefix, **h} for h in result.rows]
            for row in rows:
                row["sector_source"] = result.source if row["sector"] else ""
            return summary, rows

        summary.update(status="no_data",
                       detail="; ".join(failures) or "no source covers this ETF")
        return summary, []


# ---------------------------------------------------------------------- Sectors


def _ticker_key(ticker: str) -> str:
    """'BRK.B', 'BRK/B' and 'BRKB' all match."""
    return re.sub(r"[^A-Z0-9]", "", ticker.upper())


class SectorLookup:
    """Sectors of securities as iShares classifies them, to fill sectors other sources leave blank.

    Zacks publishes no sectors and State Street's files leave them as "-". Using a
    single issuer's classification keeps one taxonomy (iShares uses GICS sectors).
    Securities match by ISIN or CUSIP, or by ticker for US-listed ones.
    """

    SOURCE = "ishares"

    def __init__(self):
        self.by_id: dict[str, str] = {}
        self.by_ticker: dict[str, str] = {}

    def learn(self, row: dict) -> None:
        if row["source"] != self.SOURCE or not row["sector"]:
            return
        if "equity" not in str(row["asset_class"]).lower():
            return  # cash, futures and the like are "Cash and/or Derivatives"
        for key in (row["isin"], row["cusip"]):
            if key:
                self.by_id.setdefault(key, row["sector"])
        if row["ticker"] and row["country"] in ("United States", ""):
            self.by_ticker.setdefault(_ticker_key(row["ticker"]), row["sector"])

    def lookup(self, row: dict) -> str:
        for key in (row.get("isin"), row.get("cusip")):
            if key and key in self.by_id:
                return self.by_id[key]
        ticker = row.get("ticker")
        return self.by_ticker.get(_ticker_key(ticker), "") if ticker else ""


def backfill_sectors(path: Path, lookup: SectorLookup) -> tuple[int, int]:
    """Fill blank sectors in a holdings CSV in place. Returns (filled, still blank)."""
    filled = blank = 0
    tmp = path.with_name(f"{path.name}.tmp")
    with path.open(newline="", encoding="utf-8") as src, \
            tmp.open("w", newline="", encoding="utf-8") as dst:
        writer = csv.DictWriter(dst, fieldnames=HOLDING_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        for row in csv.DictReader(src):
            if not row["sector"]:
                sector = lookup.lookup(row)
                if sector:
                    row["sector"], row["sector_source"] = sector, f"{SectorLookup.SOURCE}_lookup"
                    filled += 1
                else:
                    blank += 1
            writer.writerow(row)
    os.replace(tmp, path)
    return filled, blank


# -------------------------------------------------------------------------- Main


def parse_args(argv: list[str] | None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Scrape the latest holdings of every US-listed ETF from issuer websites.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument("--tickers", nargs="+", metavar="TICKER",
                        help="Only these tickers (space or comma separated) instead of all ETFs")
    parser.add_argument("--limit", type=int, help="Only the first N ETFs, alphabetically")
    parser.add_argument("--sources", nargs="+", choices=list(ISSUER_SOURCES),
                        default=list(ISSUER_SOURCES), help="Issuer sources to use")
    parser.add_argument("--no-fallback", action="store_true",
                        help="Skip Zacks: only ETFs covered by an issuer source get holdings")
    parser.add_argument("--browser", default="chrome",
                        help='Browser for Zacks pages: "chrome" (your installed Google Chrome), '
                             '"chromium" (after `playwright install chromium`), or a path')
    parser.add_argument("--browser-workers", type=int, default=4,
                        help="Browsers loading Zacks pages in parallel")
    parser.add_argument("--show-browser", action="store_true",
                        help="Show the browser windows instead of running them in the background")
    parser.add_argument("--output-dir", type=Path, default=Path("output"))
    parser.add_argument("--cache-dir", type=Path, default=Path(".cache"),
                        help="Downloads are kept here for a few hours so reruns resume quickly")
    parser.add_argument("--no-cache", action="store_true", help="Do not read or write the cache")
    parser.add_argument("--refresh", action="store_true", help="Ignore cached downloads")
    parser.add_argument("--workers", type=int, default=8, help="Parallel download threads")
    parser.add_argument("--rate", type=float, default=2.0,
                        help="Max requests per second to any one website")
    parser.add_argument("--user-agent", default=DEFAULT_USER_AGENT)
    parser.add_argument("-v", "--verbose", action="store_true")
    args = parser.parse_args(argv)

    if args.rate <= 0:
        parser.error("--rate must be positive")
    if args.workers < 1 or args.browser_workers < 1:
        parser.error("--workers and --browser-workers must be at least 1")
    if args.tickers:
        args.tickers = [
            t.strip().upper() for arg in args.tickers for t in arg.split(",") if t.strip()]
    return args


def select_etfs(universe: list[Etf], tickers: list[str] | None, limit: int | None) -> list[Etf]:
    etfs = universe
    if tickers:
        by_ticker = {e.ticker: e for e in universe}
        missing = [t for t in tickers if t not in by_ticker]
        if missing:
            log.warning("Not in the US ETF list, trying anyway: %s", ", ".join(missing))
        etfs = [by_ticker.get(t, Etf(t, "", "")) for t in dict.fromkeys(tickers)]
    if limit:
        etfs = etfs[:limit]
    return etfs


def write_csv(path: Path, columns: list[str], rows: Iterable[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def run(args: argparse.Namespace, http: HttpClient | None = None,
        browser: BrowserPool | None = None) -> int:
    http = http or HttpClient(args.user_agent, args.rate)
    cache = DiskCache(None if args.no_cache else args.cache_dir)
    fetcher = Fetcher(http, cache, args.refresh)
    args.output_dir.mkdir(parents=True, exist_ok=True)

    universe = load_universe(fetcher)
    log.info("Found %d US-listed ETFs", len(universe))
    etfs = select_etfs(universe, args.tickers, args.limit)

    issuers = [ISSUER_SOURCES[name](fetcher) for name in args.sources]
    for source in issuers:
        source.prepare()
    fallback = None
    if not args.no_fallback:
        # Zacks turns away plain HTTP clients, so its pages are loaded in a browser.
        browser = browser or BrowserPool(args.browser_workers, args.rate, args.browser,
                                         headless=not args.show_browser)
        browser.start()
        fallback = ZacksSource(Fetcher(browser, cache, args.refresh))
    try:
        return scrape_all(args, etfs, HoldingsScraper(issuers, fallback))
    finally:
        if browser is not None:
            browser.close()


def scrape_all(args: argparse.Namespace, etfs: list[Etf], scraper: HoldingsScraper) -> int:
    planned = Counter(scraper.sources_for(etf)[0].name if scraper.sources_for(etf) else "none"
                      for etf in etfs)
    log.info("Primary source per ETF: %s",
             ", ".join(f"{k}={v}" for k, v in planned.most_common()))

    holdings_path = args.output_dir / "holdings.csv"
    summary_path = args.output_dir / "etf_summary.csv"
    summaries: list[dict] = []
    statuses: Counter = Counter()
    sectors = SectorLookup()
    holdings_total = 0
    interrupted = False
    started = time.monotonic()

    with holdings_path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=HOLDING_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        pool = ThreadPoolExecutor(max_workers=args.workers)
        try:
            futures = [pool.submit(scraper.scrape, etf) for etf in etfs]
            for done, future in enumerate(as_completed(futures), 1):
                summary, rows = future.result()
                writer.writerows(rows)  # only this thread writes, so no lock is needed
                for row in rows:
                    sectors.learn(row)
                summaries.append(summary)
                statuses[summary["source"] or summary["status"]] += 1
                holdings_total += len(rows)
                if done % 50 == 0 or done == len(futures):
                    log.info("[%d/%d] %s | %d holdings | %.0fs", done, len(futures),
                             " ".join(f"{k}={v}" for k, v in sorted(statuses.items())),
                             holdings_total, time.monotonic() - started)
        except KeyboardInterrupt:
            interrupted = True
            log.warning("Interrupted; writing what has been collected so far")
            pool.shutdown(wait=False, cancel_futures=True)
        else:
            pool.shutdown()

    filled, blank = backfill_sectors(holdings_path, sectors)
    log.info("Filled in sectors from iShares for %d holdings; %d still have none", filled, blank)
    summaries.sort(key=lambda s: s["ticker"])
    write_csv(summary_path, SUMMARY_COLUMNS, summaries)
    found = sum(1 for s in summaries if s["status"] == "ok")
    log.info("Wrote %d holdings for %d of %d ETFs to %s",
             holdings_total, found, len(summaries), holdings_path)
    log.info("Wrote per-ETF status to %s", summary_path)
    return 130 if interrupted else 0


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        datefmt="%H:%M:%S",
    )
    if not args.verbose:
        logging.getLogger("urllib3").setLevel(logging.WARNING)
    try:
        return run(args)
    except requests.RequestException as exc:  # per-ETF failures are handled inside run()
        log.error("Could not download the ETF list: %s", exc)
        return 1
    except SetupError as exc:
        log.error("Zacks fallback unavailable: %s. Install Google Chrome, or run "
                  "`playwright install chromium` and pass --browser chromium, or pass "
                  "--no-fallback to use issuer websites only.", exc)
        return 1


if __name__ == "__main__":
    sys.exit(main())
